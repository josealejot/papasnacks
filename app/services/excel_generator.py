from io import BytesIO
from datetime import datetime
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.config import settings
from app.models.cliente import Cliente
from app.models.producto import Producto
from app.models.venta import Venta
from app.models.transaccion import Transaccion
from app.models.gasto import Gasto
from app.services.capital_service import obtener_resumen_capital


CURRENCY_FORMAT = '"$" #,##0'


def _estilizar_hoja(ws, titulo: str, subtitulo: str, encabezados: list[str], color_header: str = "1E293B"):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(encabezados))
    celda_titulo = ws.cell(row=1, column=1, value=titulo)
    celda_titulo.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    celda_titulo.fill = PatternFill(start_color=color_header, end_color=color_header, fill_type="solid")
    celda_titulo.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(encabezados))
    celda_sub = ws.cell(row=2, column=1, value=subtitulo)
    celda_sub.font = Font(name="Calibri", size=10, italic=True, color="475569")
    celda_sub.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ws.row_dimensions[2].height = 20

    borde_fino = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws.row_dimensions[4].height = 24
    for col_idx, texto in enumerate(encabezados, start=1):
        c = ws.cell(row=4, column=col_idx, value=texto)
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borde_fino


def _autoajustar_columnas(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in (1, 2):
                continue
            val_str = str(cell.value or "")
            if cell.number_format == CURRENCY_FORMAT and isinstance(cell.value, (int, float)):
                val_str = f"$ {cell.value:,.0f}"
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(min(max_len + 5, 42), 13)


def generar_excel_completo(db: Session) -> BytesIO:
    """Genera un archivo Excel (.xlsx) multipágina con todo el control financiero y comercial."""
    wb = Workbook()
    fecha_gen = datetime.now().strftime("%d/%m/%Y %H:%M")
    resumen = obtener_resumen_capital(db)

    borde_celda = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    fill_verde = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_amarillo = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
    fill_rojo = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    # =========================================================================
    # HOJA 1: RESUMEN DE CAPITAL Y UTILIDADES
    # =========================================================================
    ws_cap = wb.active
    ws_cap.title = "Resumen de Capital"
    _estilizar_hoja(
        ws_cap,
        f"{settings.BUSINESS_NAME} (NIT: {settings.BUSINESS_NIT}) - RESUMEN EJECUTIVO DE CAPITAL Y UTILIDADES",
        f"Corte generado el {fecha_gen} | Tel: {settings.BUSINESS_PHONE} | Base: Costo $2.700 -> Venta $4.000",
        ["Indicador Financiero", "Categoría", "Valor (COP)", "Interpretación Operativa"],
        color_header="0F172A",
    )

    metricas = [
        ("Dinero en Efectivo (Caja Física)", "Liquidez Disponible", resumen["caja_efectivo"], "Efectivo real disponible en caja"),
        ("Dinero en Transferencias (Bancos / Nequi / Daviplata)", "Liquidez Disponible", resumen["bancos_transferencia"], "Saldo disponible en cuentas digitales"),
        ("LIQUIDEZ TOTAL DISPONIBLE (Efectivo + Transferencias)", "Liquidez Total", resumen["liquidez_total"], "Dinero total en mano listo para usar"),
        ("Capital en la Calle (Cartera por Cobrar)", "Cuentas por Cobrar", resumen["capital_en_la_calle"], "Saldo pendiente que deben los clientes"),
        ("Capital en Inventario (Valorizado al Costo $2.700)", "Inventario", resumen["capital_inventario_costo"], f"{resumen['stock_total_unidades']} unidades en stock al costo"),
        ("Valor de Venta Proyectado del Inventario ($4.000)", "Inventario", resumen["capital_inventario_venta"], "Ingreso proyectado al vender el stock actual"),
        ("PATRIMONIO / CAPITAL TOTAL DEL NEGOCIO", "Patrimonio Neto", resumen["patrimonio_total"], "Liquidez + Cartera + Inventario al Costo"),
        ("Ventas Totales Facturadas", "Resultados Comerciales", resumen["ventas_totales"], f"{resumen['unidades_vendidas']} unidades vendidas"),
        ("Costo Total de Mercancía Vendida", "Costos Directos", resumen["costo_mercancia_vendida"], "Costo proveedor de las unidades vendidas"),
        ("Utilidad Bruta en Ventas", "Rentabilidad", resumen["utilidad_bruta_ventas"], f"Margen bruto del {resumen['margen_bruto_pct']}%"),
        ("Utilidad Realizada (Efectivamente Cobrada)", "Rentabilidad", resumen["utilidad_realizada"], "Ganancia ya recaudada en Caja/Transferencia"),
        ("Gastos Operativos (Fletes, Empaques, Otros)", "Gastos", resumen["gastos_operativos"], "Gastos operativos distintos a inventario"),
        ("UTILIDAD NETA REAL DEL EJERCICIO", "Utilidad Final", resumen["utilidad_neta"], f"Margen neto del {resumen['margen_neto_pct']}%"),
    ]

    for idx, (nombre, cat, valor, nota) in enumerate(metricas, start=5):
        c1 = ws_cap.cell(row=idx, column=1, value=nombre)
        c2 = ws_cap.cell(row=idx, column=2, value=cat)
        c3 = ws_cap.cell(row=idx, column=3, value=valor)
        c4 = ws_cap.cell(row=idx, column=4, value=nota)
        c3.number_format = CURRENCY_FORMAT
        for c in (c1, c2, c3, c4):
            c.border = borde_celda
        if "TOTAL" in nombre or "NETA" in nombre:
            for c in (c1, c2, c3, c4):
                c.font = Font(bold=True)
                c.fill = fill_verde if valor >= 0 else fill_rojo

    _autoajustar_columnas(ws_cap)

    # =========================================================================
    # HOJA 2: VENTAS Y FACTURACIÓN
    # =========================================================================
    ws_ven = wb.create_sheet(title="Ventas y Facturas")
    encabezados_ven = [
        "N° Factura",
        "Fecha",
        "Cliente",
        "Unidades",
        "Costo Total",
        "Total Venta",
        "Utilidad Bruta",
        "Pagado Efectivo",
        "Pagado Transferencia",
        "Total Abonado",
        "Saldo Pendiente",
        "Estado",
    ]
    _estilizar_hoja(
        ws_ven,
        "HISTORIAL DETALLADO DE VENTAS Y FACTURACIÓN",
        f"Generado el {fecha_gen}",
        encabezados_ven,
        color_header="1E3A8A",
    )

    ventas = db.query(Venta).filter(Venta.estado != "ANULADA").order_by(Venta.fecha.desc()).all()
    fila = 5
    for v in ventas:
        datos = [
            v.numero_factura,
            v.fecha.strftime("%d/%m/%Y %H:%M"),
            v.cliente.nombre if v.cliente else "Cliente General",
            v.total_unidades,
            v.subtotal_costo,
            v.total_venta,
            v.utilidad_bruta,
            v.pagado_efectivo,
            v.pagado_transferencia,
            v.monto_pagado,
            v.saldo_pendiente,
            v.estado,
        ]
        for col_idx, val in enumerate(datos, start=1):
            celda = ws_ven.cell(row=fila, column=col_idx, value=val)
            celda.border = borde_celda
            if col_idx in (5, 6, 7, 8, 9, 10, 11):
                celda.number_format = CURRENCY_FORMAT
            if col_idx == 12:
                celda.alignment = Alignment(horizontal="center")
                celda.font = Font(bold=True)
                if val == "PAGADO":
                    celda.fill = fill_verde
                elif val == "PARCIAL":
                    celda.fill = fill_amarillo
                else:
                    celda.fill = fill_rojo
        fila += 1

    # Fila de Totales
    ws_ven.cell(row=fila, column=1, value="TOTALES").font = Font(bold=True)
    ws_ven.cell(row=fila, column=4, value=sum(v.total_unidades for v in ventas)).font = Font(bold=True)
    totales_cols = {
        5: sum(v.subtotal_costo for v in ventas),
        6: sum(v.total_venta for v in ventas),
        7: sum(v.utilidad_bruta for v in ventas),
        8: sum(v.pagado_efectivo for v in ventas),
        9: sum(v.pagado_transferencia for v in ventas),
        10: sum(v.monto_pagado for v in ventas),
        11: sum(v.saldo_pendiente for v in ventas),
    }
    for col_idx in range(1, 13):
        c = ws_ven.cell(row=fila, column=col_idx)
        c.fill = fill_total
        c.border = borde_celda
        if col_idx in totales_cols:
            c.value = totales_cols[col_idx]
            c.number_format = CURRENCY_FORMAT
            c.font = Font(bold=True)

    _autoajustar_columnas(ws_ven)

    # =========================================================================
    # HOJA 3: CARTERA POR COBRAR (DEUDAS PENDIENTES)
    # =========================================================================
    ws_car = wb.create_sheet(title="Cartera por Cobrar")
    encabezados_car = [
        "N° Factura",
        "Fecha Venta",
        "Cliente",
        "Teléfono",
        "Total Factura",
        "Abonado Efectivo",
        "Abonado Transferencia",
        "Total Abonado",
        "Saldo Pendiente",
        "Estado",
    ]
    _estilizar_hoja(
        ws_car,
        "CONTROL DE CARTERA - CAPITAL EN LA CALLE POR COBRAR",
        f"Facturas con saldo pendiente al {fecha_gen}",
        encabezados_car,
        color_header="991B1B",
    )

    pendientes = [v for v in ventas if v.saldo_pendiente > 0]
    fila = 5
    for v in pendientes:
        datos = [
            v.numero_factura,
            v.fecha.strftime("%d/%m/%Y"),
            v.cliente.nombre if v.cliente else "",
            v.cliente.telefono if v.cliente else "",
            v.total_venta,
            v.pagado_efectivo,
            v.pagado_transferencia,
            v.monto_pagado,
            v.saldo_pendiente,
            v.estado,
        ]
        for col_idx, val in enumerate(datos, start=1):
            celda = ws_car.cell(row=fila, column=col_idx, value=val)
            celda.border = borde_celda
            if col_idx in (5, 6, 7, 8, 9):
                celda.number_format = CURRENCY_FORMAT
            if col_idx == 9:
                celda.font = Font(bold=True, color="991B1B")
                celda.fill = fill_rojo
            if col_idx == 10:
                celda.fill = fill_amarillo if val == "PARCIAL" else fill_rojo
        fila += 1

    ws_car.cell(row=fila, column=1, value="TOTAL POR COBRAR").font = Font(bold=True)
    c_tot_car = ws_car.cell(row=fila, column=9, value=sum(v.saldo_pendiente for v in pendientes))
    c_tot_car.number_format = CURRENCY_FORMAT
    c_tot_car.font = Font(bold=True, color="991B1B")
    c_tot_car.fill = fill_rojo
    _autoajustar_columnas(ws_car)

    # =========================================================================
    # HOJA 4: INVENTARIO Y RENTABILIDAD POR PRODUCTO
    # =========================================================================
    ws_inv = wb.create_sheet(title="Inventario Snacks")
    encabezados_inv = [
        "Código",
        "Producto / Snack",
        "Sabor",
        "Presentación",
        "Costo Compra",
        "Precio Venta",
        "Ganancia x Unidad",
        "Stock Actual",
        "Capital al Costo",
        "Valor en Venta",
        "Utilidad Potencial",
    ]
    _estilizar_hoja(
        ws_inv,
        "INVENTARIO VALORIZADO DE SNACKS",
        f"Generado el {fecha_gen}",
        encabezados_inv,
        color_header="065F46",
    )

    productos = db.query(Producto).filter(Producto.activo == True).order_by(Producto.codigo.asc()).all()
    fila = 5
    for p in productos:
        datos = [
            p.codigo,
            p.nombre,
            p.sabor,
            p.presentacion,
            p.costo_compra,
            p.precio_venta,
            p.ganancia_unitaria,
            p.stock_actual,
            p.valor_inventario_costo,
            p.valor_inventario_venta,
            p.utilidad_potencial_stock,
        ]
        for col_idx, val in enumerate(datos, start=1):
            celda = ws_inv.cell(row=fila, column=col_idx, value=val)
            celda.border = borde_celda
            if col_idx in (5, 6, 7, 9, 10, 11):
                celda.number_format = CURRENCY_FORMAT
        fila += 1

    _autoajustar_columnas(ws_inv)

    # =========================================================================
    # HOJA 5: MOVIMIENTOS DE CAJA Y TRANSFERENCIAS
    # =========================================================================
    ws_tx = wb.create_sheet(title="Flujo Caja y Transferencias")
    encabezados_tx = [
        "Recibo",
        "Fecha y Hora",
        "Tipo Movimiento",
        "Canal",
        "Efectivo",
        "Transferencia",
        "Total Movimiento",
        "Concepto / Detalle",
    ]
    _estilizar_hoja(
        ws_tx,
        "LIBRO DIARIO DE CAJA (EFECTIVO) Y BANCOS (TRANSFERENCIAS)",
        f"Generado el {fecha_gen}",
        encabezados_tx,
        color_header="312E81",
    )

    txs = db.query(Transaccion).filter(Transaccion.anulada == False).order_by(Transaccion.fecha.desc()).all()
    fila = 5
    for t in txs:
        signo = -1 if t.es_egreso else 1
        datos = [
            t.codigo_recibo or f"TX-{t.id}",
            t.fecha.strftime("%d/%m/%Y %H:%M"),
            t.tipo,
            t.canal,
            t.monto_efectivo * signo,
            t.monto_transferencia * signo,
            t.monto_total * signo,
            t.concepto,
        ]
        for col_idx, val in enumerate(datos, start=1):
            celda = ws_tx.cell(row=fila, column=col_idx, value=val)
            celda.border = borde_celda
            if col_idx in (5, 6, 7):
                celda.number_format = CURRENCY_FORMAT
                celda.fill = fill_verde if signo > 0 else fill_rojo
        fila += 1

    _autoajustar_columnas(ws_tx)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
